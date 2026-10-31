"""COMPASS I/O tests"""

import os
import multiprocessing
from pathlib import Path
from unittest.mock import AsyncMock

import pytest

pytest.importorskip("docling_rs")
pytest.importorskip("docling_core")

from docx import Document
from elm.web.search.run import load_docs
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont

import compass.scripts.download as download_module
import compass.web.file_loader as file_loader_module
from compass.pb import COMPASS_PB
from compass.pipeline.collection.dedupe import DocumentDeDuplicator
from compass.pipeline.collection.persistence import (
    load_collected_docs,
    persist_documents,
)
from compass.services.docling import (
    FileLoader,
    read_docling_local_file,
    read_docling_web_file,
    _run_docling_in_subprocess,
)
from compass.services.threaded import (
    OrdFileMover,
    ParsedFileWriter,
    TempFileCache,
    TempFileCacheCopier,
    TempFileCachePB,
)
from compass.web.file_loader import (
    AsyncLocalDoclingFileLoader,
    AsyncDoclingWebFileLoader,
)
from compass.services.provider import RunningAsyncServices
from compass.exceptions import COMPASSNotInitializedError
from compass.utilities.jurisdictions import Jurisdiction
from compass.utilities.logs import LogListener
from compass.utilities.parsing import is_pdf_doc
from compass.warn import COMPASSWarning


PYT_CMD = os.getenv("TESSERACT_CMD")


def _return_from_subprocess(value):
    """Return a serializable value from a child process"""
    return value


@pytest.fixture(autouse=True)
async def subprocess_log_listener():
    """Drain worker logs until process-pool shutdown completes"""
    async with LogListener([]):
        yield


@pytest.fixture
def production_docling_loaders():
    """Require Docling selection before production modules are imported"""
    if os.getenv("COMPASS_FILE_LOAD_BACKEND") != "docling":
        pytest.skip("Requires COMPASS_FILE_LOAD_BACKEND=docling")
    assert (
        file_loader_module.COMPASSLocalFileLoader
        is download_module.COMPASSLocalFileLoader
        is AsyncLocalDoclingFileLoader
    )
    assert (
        file_loader_module.COMPASSWebFileLoader
        is download_module.COMPASSWebFileLoader
        is AsyncDoclingWebFileLoader
    )
    return (
        download_module.COMPASSLocalFileLoader,
        download_module.COMPASSWebFileLoader,
    )


@pytest.fixture
def scanned_ordinance_pdf(tmp_path):
    """Create a compact scanned ordinance with known setback values"""
    source = tmp_path / "scanned_ordinance.pdf"
    image = Image.new("RGB", (1600, 900), "white")
    drawing = ImageDraw.Draw(image)
    drawing.multiline_text(
        (80, 80),
        "EXAMPLE COUNTY WIND ORDINANCE\n\n"
        "Road setback: 125 feet\n\n"
        "Building setback: 250 feet",
        fill="black",
        font=ImageFont.load_default(size=36),
        spacing=24,
    )
    image.save(source, format="PDF", resolution=150)
    image.close()
    return source


@pytest.fixture
def small_ordinance(tmp_path):
    """Create compact ordinance fixtures without model downloads"""
    markdown = (
        "# Example County Wind Ordinance\n\n"
        "Wind turbines require a minimum setback of 125 feet.\n\n"
        "| Structure | Setback |\n"
        "| --- | --- |\n"
        "| Roads | 125 feet |\n"
        "| Buildings | 250 feet |\n"
    )
    html = (
        "<!DOCTYPE html><html><body>"
        "<h1>Example County Wind Ordinance</h1>"
        "<p>Wind turbines require a minimum setback of 125 feet.</p>"
        "<table><tr><th>Structure</th><th>Setback</th></tr>"
        "<tr><td>Roads</td><td>125 feet</td></tr>"
        "<tr><td>Buildings</td><td>250 feet</td></tr></table>"
        '<a href="appendix.html">Ordinance appendix</a>'
        "</body></html>"
    )
    files = {}
    for extension, content in (
        ("md", markdown),
        ("txt", markdown),
        ("html", html),
    ):
        fp = tmp_path / f"ordinance.{extension}"
        fp.write_text(content, encoding="utf-8")
        files[extension] = fp
    return files


@pytest.mark.asyncio
@pytest.mark.parametrize("extension", ["csv", "docx", "xlsx"])
async def test_real_docling_office_conversion_preserves_table(
    tmp_path, extension
):
    """Read supported office formats through the shared converter path"""
    source = tmp_path / f"setbacks.{extension}"
    rows = [
        ["Structure", "Setback"],
        ["Roads", "125 feet"],
        ["Buildings", "250 feet"],
    ]
    if extension == "docx":
        document = Document()
        table = document.add_table(rows=0, cols=2)
        for row in rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, row, strict=True):
                cell.text = value
        document.save(source)
    elif extension == "xlsx":
        workbook = Workbook()
        for row in rows:
            workbook.active.append(row)
        workbook.save(source)
        workbook.close()
    else:
        source.write_text(
            "Structure,Setback\nRoads,125 feet\nBuildings,250 feet\n",
            encoding="utf-8",
        )

    async with RunningAsyncServices([FileLoader()]):
        doc, raw = await read_docling_local_file(
            source, pdf_pipeline_options={"document_timeout": 30}
        )

    assert not doc.empty
    assert doc.attrs["doc_type"] == extension
    assert doc.attrs["conversion_status"] == "success"
    assert raw == source.read_bytes()
    for expected in (
        "Structure",
        "Setback",
        "Roads",
        "Buildings",
        "125",
        "250",
    ):
        assert expected in doc.text


@pytest.mark.asyncio
async def test_real_docling_ocr_preserves_known_scanned_text(
    scanned_ordinance_pdf,
):
    """Recover known setback values from a generated image-only PDF"""
    source = scanned_ordinance_pdf
    async with RunningAsyncServices([FileLoader()]):
        doc, raw = await read_docling_local_file(
            source,
            pytesseract_exe_fp=PYT_CMD,
            pdf_pipeline_options={
                "do_table_structure": False,
                "document_timeout": 120,
            },
        )

    assert not doc.empty
    assert doc.attrs["from_ocr"] is True
    assert doc.attrs["num_pages"] == 1
    assert doc.attrs["conversion_status"] == "success"
    assert "EXAMPLE COUNTY WIND ORDINANCE" in doc.text.upper()
    assert "125" in doc.text
    assert "250" in doc.text
    assert raw == source.read_bytes()


@pytest.mark.docling_backend
def test_production_docling_backend_disables_elm_fallback(
    production_docling_loaders,
):
    """Use production Docling aliases without an ELM conversion fallback"""
    __, web_loader = production_docling_loaders
    assert web_loader(re_fetch_failed_with_elm=False).failed_fetcher is None


@pytest.mark.docling_backend
@pytest.mark.asyncio
@pytest.mark.parametrize("source_kind", ["text", "scan"])
@pytest.mark.parametrize("entry_point", ["local", "web"])
async def test_production_docling_collection_roundtrip(
    production_docling_loaders,
    test_data_files_dir,
    scanned_ordinance_pdf,
    tmp_path,
    monkeypatch,
    source_kind,
    entry_point,
):
    """Persist and reload real PDFs through Docling collection entry points"""
    source = (
        scanned_ordinance_pdf
        if source_kind == "scan"
        else test_data_files_dir / "Caneadea New York.pdf"
    )
    raw = source.read_bytes()
    jurisdiction = Jurisdiction("county", "Colorado", county="Example")
    url = "https://example.com/ordinance.pdf"
    options = {
        "do_ocr": source_kind == "scan",
        "do_table_structure": False,
        "document_timeout": 120,
    }
    monkeypatch.setattr(
        file_loader_module.AsyncFetchWithRetry,
        "fetch",
        AsyncMock(return_value=(raw, "application/pdf", None, {})),
    )
    services = [
        FileLoader(),
        TempFileCachePB(),
        OrdFileMover(tmp_path),
        ParsedFileWriter(tmp_path),
        TempFileCacheCopier(),
    ]
    COMPASS_PB.reset()
    try:
        async with RunningAsyncServices(services):
            with COMPASS_PB.jurisdiction_prog_bar(
                jurisdiction.full_name, progress_main=False
            ):
                if entry_point == "local":
                    docs = await download_module.load_known_docs(
                        jurisdiction,
                        [source],
                        local_file_loader_kwargs={
                            "pdf_pipeline_options": options,
                            "doc_attrs": {"user_label": "test ordinance"},
                        },
                    )
                else:
                    docs = await download_module.download_known_urls(
                        jurisdiction,
                        [url],
                        file_loader_kwargs={
                            "pdf_pipeline_options": options,
                            "re_fetch_failed_with_elm": False,
                        },
                    )

            assert len(docs) == 1
            doc = docs[0]
            assert not doc.empty
            assert is_pdf_doc(doc)
            assert doc.attrs["doc_type"] == "pdf"
            assert doc.attrs["conversion_status"] == "success"
            assert doc.attrs["from_ocr"] is (source_kind == "scan")
            assert doc.attrs["num_pages"] == (
                1 if source_kind == "scan" else 3
            )
            if source_kind == "scan":
                assert "EXAMPLE COUNTY WIND ORDINANCE" in doc.text.upper()
                assert "125" in doc.text
                assert "250" in doc.text
            else:
                assert "Town of Caneadea" in doc.text
                assert "Section 15.4" in doc.text
            if entry_point == "local":
                assert doc.attrs["source_fp"] == source
                assert doc.attrs["user_label"] == "test ordinance"
            else:
                assert doc.attrs["source"] == url
            assert Path(doc.attrs["cache_fn"]).read_bytes() == raw
            parsed_text = doc.text
            checksum = doc.attrs["checksum"]

            collected = DocumentDeDuplicator()
            collected.add_docs(docs, step_name=entry_point)
            collected.add_docs(docs, step_name="repeat")
            assert len(collected) == 1
            collection_info = await persist_documents(
                jurisdiction, collected, [entry_point, "repeat"]
            )
            assert collection_info["num_docs"] == 1
            assert collection_info["completed_step_document_counts"] == {
                entry_point: 1,
                "repeat": 1,
            }
            record = collection_info["results"][0]
            assert record["is_pdf"] is True
            assert record["num_pages"] == doc.attrs["num_pages"]
            assert record["checksum"] == checksum
            assert record["from_steps"] == [entry_point, "repeat"]
            assert "cache_fn" not in record
            assert Path(record["source_fp"]).read_bytes() == raw
            assert (
                Path(record["parsed_fp"]).read_text(encoding="utf-8")
                == parsed_text
            )

            reloaded = await load_collected_docs(
                collection_info, task_name=jurisdiction.full_name
            )
            assert len(reloaded) == 1
            restored = reloaded[0]
            assert not restored.empty
            assert is_pdf_doc(restored)
            assert restored.text.split() == parsed_text.split()
            assert not restored.remove_comments
            for key, value in record.items():
                assert restored.attrs[key] == value
            assert Path(restored.attrs["cache_fn"]).read_bytes() == raw
    finally:
        COMPASS_PB.reset()


@pytest.mark.docling_backend
@pytest.mark.asyncio
async def test_production_docling_corrupt_pdf_has_no_elm_fallback(
    production_docling_loaders, monkeypatch, tmp_path
):
    """Reject corrupt downloaded PDFs without caching or ELM retries"""
    __, web_loader = production_docling_loaders
    loader = web_loader(
        re_fetch_failed_with_elm=False,
        file_cache_coroutine=TempFileCache.call,
        pdf_pipeline_options={"document_timeout": 30},
    )
    fetch = AsyncMock(return_value=(b"not a PDF", "application/pdf", None, {}))
    monkeypatch.setattr(loader.content_fetcher, "fetch", fetch)
    cache = TempFileCache(td_kwargs={"dir": tmp_path})
    async with RunningAsyncServices([FileLoader(), cache]):
        assert await loader.fetch_all("https://example.com/corrupt.pdf") == []
        assert not any(path.is_file() for path in tmp_path.rglob("*"))
    assert loader.failed_fetcher is None
    fetch.assert_awaited_once_with("https://example.com/corrupt.pdf")


@pytest.mark.asyncio
@pytest.mark.parametrize("extension", ["md", "txt", "html"])
async def test_real_docling_local_conversion_and_cache(
    small_ordinance, extension
):
    """Preserve real headings, table values, metadata, and cached bytes"""
    source = small_ordinance[extension]
    loader = AsyncLocalDoclingFileLoader(
        doc_attrs={"source": "https://example.com/ordinance"},
        file_cache_coroutine=TempFileCache.call,
        pdf_pipeline_options={"document_timeout": 30},
    )
    async with RunningAsyncServices([FileLoader(), TempFileCache()]):
        doc = await loader.fetch(source)

        assert not doc.empty
        for expected in (
            "Example County Wind Ordinance",
            "minimum setback",
            "Roads",
            "Buildings",
            "125 feet",
            "250 feet",
        ):
            assert expected in doc.text
        assert "|" in doc.text
        assert doc.attrs["doc_type"] == (
            "html" if extension == "html" else "md"
        )
        assert doc.attrs["conversion_status"] == "success"
        assert doc.attrs["source_fp"] == source
        assert doc.attrs["source"] == "https://example.com/ordinance"
        assert doc.attrs["from_ocr"]
        cache_fp = Path(doc.attrs["cache_fn"])
        if extension == "html":
            assert cache_fp.read_text(encoding="utf-8") == doc.text
        else:
            assert cache_fp.read_bytes() == source.read_bytes()
        assert doc.attrs["checksum"]


@pytest.mark.asyncio
async def test_real_docling_web_html_conversion(small_ordinance, monkeypatch):
    """Parse downloaded HTML with its original URI and real converter"""
    raw = small_ordinance["html"].read_bytes()
    url = "https://example.com/code/ordinance"
    loader = AsyncDoclingWebFileLoader(
        file_cache_coroutine=TempFileCache.call,
        pdf_pipeline_options={"document_timeout": 30},
    )
    monkeypatch.setattr(
        loader.content_fetcher,
        "fetch",
        AsyncMock(
            return_value=(
                raw,
                "text/html",
                "utf-8",
                {
                    "Content-Type": "text/html",
                    "Content-Disposition": 'inline; filename="ordinance.html"',
                },
            )
        ),
    )
    async with RunningAsyncServices([FileLoader(), TempFileCache()]):
        doc = await loader.fetch(url)

        assert not doc.empty
        assert "Example County Wind Ordinance" in doc.text
        assert "250 feet" in doc.text
        assert "https://example.com/code/appendix.html" in doc.text
        assert doc.attrs["source"] == url
        assert doc.attrs["doc_filename"] == "ordinance"
        assert doc.attrs["doc_type"] == "html"
        assert (
            Path(doc.attrs["cache_fn"]).read_text(encoding="utf-8") == doc.text
        )


@pytest.mark.asyncio
async def test_real_docling_invalid_pdf_returns_empty_document():
    """Handle corrupt downloaded PDFs without downloading model assets"""
    async with RunningAsyncServices([FileLoader()]):
        doc = await read_docling_web_file(
            b"not a PDF",
            "corrupt.pdf",
            pdf_pipeline_options={"document_timeout": 30},
        )

    assert doc.empty
    assert doc.attrs["doc_type"] == "unknown"


@pytest.mark.asyncio
async def test_real_docling_collection_reload_preserves_artifacts(
    small_ordinance,
):
    """Reload persisted parsed text and cache its unchanged source file"""
    doc_info = {
        "parsed_fp": small_ordinance["txt"],
        "source_fp": small_ordinance["html"],
        "source": "https://example.com/ordinance",
        "jurisdiction": "Example County",
        "from_steps": ["search"],
    }
    async with RunningAsyncServices([FileLoader(), TempFileCacheCopier()]):
        docs = await load_collected_docs(
            {"results": [doc_info]}, task_name="Example County"
        )

        assert len(docs) == 1
        doc = docs[0]
        assert not doc.empty
        assert "Example County Wind Ordinance" in doc.text
        assert "250 feet" in doc.text
        assert not doc.remove_comments
        for key, value in doc_info.items():
            assert doc.attrs[key] == value
        cache_fp = Path(doc.attrs["cache_fn"])
        assert cache_fp != small_ordinance["html"]
        assert cache_fp.read_bytes() == small_ordinance["html"].read_bytes()


@pytest.mark.asyncio
async def test_docling_collection_reload_without_parsed_path_warns():
    """Warn and skip records that cannot be re-parsed from persisted text"""
    assert await load_collected_docs({}, task_name="Example County") == []
    with pytest.warns(COMPASSWarning, match="Parsed file path"):
        docs = await load_collected_docs(
            {"results": [{"source": "https://example.com/ordinance"}]},
            task_name="Example County",
        )
    assert docs == [None]


@pytest.mark.asyncio
async def test_basic_load_pdf(test_data_files_dir):
    """Test basic loading of local PDF document"""
    test_fp = test_data_files_dir / "Caneadea New York.pdf"

    fl = AsyncLocalDoclingFileLoader()
    async with RunningAsyncServices([FileLoader()]):
        docs = await load_docs([test_fp], fl)

    assert len(docs) == 1
    doc = docs[0]
    assert not doc.empty
    assert Path(doc.attrs.get("source_fp")) == test_fp
    assert len(doc.pages) == 1
    assert doc.attrs["num_pages"] == 3
    assert doc.attrs["doc_type"] == "pdf"
    assert "Town of Caneadea" in doc.text
    assert "Section 15.4" in doc.text


@pytest.mark.asyncio
async def test_basic_load_html(test_data_files_dir):
    """Test basic loading of local HTML document"""
    test_fp = test_data_files_dir / "Whatcom.txt"

    fl = AsyncLocalDoclingFileLoader()
    async with RunningAsyncServices([FileLoader()]):
        docs = await load_docs([test_fp], fl)

    assert len(docs) == 1
    doc = docs[0]
    assert not doc.empty
    assert Path(doc.attrs.get("source_fp")) == test_fp
    assert len(doc.pages) == 1


@pytest.mark.asyncio
async def test_basic_load_pdf_with_service(test_data_files_dir):
    """Test basic loading of local PDF document with service"""
    test_fp = test_data_files_dir / "Caneadea New York.pdf"

    with pytest.raises(
        COMPASSNotInitializedError,
        match=r"Must initialize the queue for 'FileLoader'.",
    ):
        await read_docling_local_file(test_fp)

    fl = AsyncLocalDoclingFileLoader()
    async with RunningAsyncServices([FileLoader()]):
        doc, __ = await read_docling_local_file(test_fp)
        doc_2 = await load_docs([test_fp], fl)

    assert not doc.empty
    assert not doc_2[0].empty
    assert doc.text == doc_2[0].text


@pytest.mark.asyncio
async def test_basic_load_ocr_pdf_with_service(test_data_files_dir):
    """Test basic loading of local PDF document with service"""
    test_fp = test_data_files_dir / "Sedgwick Kansas.pdf"

    with pytest.raises(
        COMPASSNotInitializedError,
        match=r"Must initialize the queue for 'FileLoader'.",
    ):
        await read_docling_local_file(test_fp)

    fl = AsyncLocalDoclingFileLoader(pytesseract_exe_fp=PYT_CMD)
    async with RunningAsyncServices([FileLoader()]):
        doc, raw = await read_docling_local_file(
            test_fp, pytesseract_exe_fp=PYT_CMD
        )
        doc_2 = await load_docs([test_fp], fl)

    assert not doc.empty
    assert not doc_2[0].empty
    assert doc.text == doc_2[0].text
    assert doc.attrs["doc_type"] == "pdf"
    assert doc.attrs["num_pages"] == 6
    assert doc.attrs["from_ocr"] is True
    assert "SEDGWICK COUNTY" in doc.text.upper()
    assert "RENEWABLE ENERGY SYSTEMS" in doc.text.upper()
    assert raw == test_fp.read_bytes()


@pytest.mark.asyncio
async def test_basic_load_html_with_service(test_data_files_dir):
    """Test basic loading of local HTML document with service"""
    test_fp = test_data_files_dir / "Whatcom.txt"

    with pytest.raises(
        COMPASSNotInitializedError,
        match=r"Must initialize the queue for 'FileLoader'.",
    ):
        await read_docling_local_file(test_fp)

    fl = AsyncLocalDoclingFileLoader()
    async with RunningAsyncServices([FileLoader()]):
        doc, __ = await read_docling_local_file(test_fp)
        doc_2 = await load_docs([test_fp], fl)

    assert not doc.empty
    assert not doc_2[0].empty
    assert doc.text == doc_2[0].text


def test_docling_subprocess_supports_spawn(monkeypatch):
    """Use spawn when fork is unavailable on the current platform"""
    monkeypatch.setattr(
        multiprocessing, "get_all_start_methods", lambda: ["spawn"]
    )
    assert (
        _run_docling_in_subprocess(
            _return_from_subprocess, args=("spawned",), kwargs={}, timeout=60
        )
        == "spawned"
    )


if __name__ == "__main__":
    pytest.main(["-q", "--show-capture=all", Path(__file__), "-rapP"])
