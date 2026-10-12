"""COMPASS I/O tests"""

import os
import multiprocessing
from pathlib import Path
from unittest.mock import AsyncMock

import pytest
from docx import Document
from elm.web.search.run import load_docs
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont

from compass.pipeline.collection.persistence import load_collected_docs
from compass.services.cpu import (
    FileLoader,
    read_docling_local_file,
    read_docling_web_file,
    _run_docling_in_subprocess,
)
from compass.services.threaded import TempFileCache, TempFileCacheCopier
from compass.web.file_loader import (
    AsyncLocalDoclingFileLoader,
    AsyncDoclingWebFileLoader,
)
from compass.services.provider import RunningAsyncServices
from compass.exceptions import COMPASSNotInitializedError
from compass.utilities.logs import LogListener
from compass.warn import COMPASSWarning


PYT_CMD = os.getenv("TESSERACT_CMD")


def _return_from_subprocess(value):
    """Return a serializable value from a child process"""
    return value


@pytest.fixture
def small_ordinance(tmp_path):
    """Create compact ordinance fixtures without model downloads"""
    markdown = (
        "# Example County Wind Ordinance\n\n"
        "Wind turbines require a minimum setback of 125 feet.\n\n"
        "| Structure | Setback |\n"
        "| --- | --- |\n"
        "| Roads | 125 feet |\n"
        "| Buildings | 250 feet |\n"
    )
    html = (
        "<!DOCTYPE html><html><body>"
        "<h1>Example County Wind Ordinance</h1>"
        "<p>Wind turbines require a minimum setback of 125 feet.</p>"
        "<table><tr><th>Structure</th><th>Setback</th></tr>"
        "<tr><td>Roads</td><td>125 feet</td></tr>"
        "<tr><td>Buildings</td><td>250 feet</td></tr></table>"
        '<a href="appendix.html">Ordinance appendix</a>'
        "</body></html>"
    )
    files = {}
    for extension, content in (
        ("md", markdown),
        ("txt", markdown),
        ("html", html),
    ):
        fp = tmp_path / f"ordinance.{extension}"
        fp.write_text(content, encoding="utf-8")
        files[extension] = fp
    return files


@pytest.mark.asyncio
@pytest.mark.parametrize("extension", ["csv", "docx", "xlsx"])
async def test_real_docling_office_conversion_preserves_table(
    tmp_path, extension
):
    """Read supported office formats through the shared converter path"""
    source = tmp_path / f"setbacks.{extension}"
    rows = [
        ["Structure", "Setback"],
        ["Roads", "125 feet"],
        ["Buildings", "250 feet"],
    ]
    if extension == "docx":
        document = Document()
        table = document.add_table(rows=0, cols=2)
        for row in rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, row, strict=True):
                cell.text = value
        document.save(source)
    elif extension == "xlsx":
        workbook = Workbook()
        for row in rows:
            workbook.active.append(row)
        workbook.save(source)
        workbook.close()
    else:
        source.write_text(
            "Structure,Setback\nRoads,125 feet\nBuildings,250 feet\n",
            encoding="utf-8",
        )

    async with RunningAsyncServices([FileLoader()]):
        doc, raw = await read_docling_local_file(
            source, pdf_pipeline_options={"document_timeout": 30}
        )

    assert not doc.empty
    assert doc.attrs["doc_type"] == extension
    assert doc.attrs["conversion_status"] == "success"
    assert raw == source.read_bytes()
    for expected in (
        "Structure",
        "Setback",
        "Roads",
        "Buildings",
        "125",
        "250",
    ):
        assert expected in doc.text


@pytest.mark.skipif(
    os.getenv("GITHUB_ACTIONS") == "true" or not PYT_CMD,
    reason="requires Tesseract and heavyweight Docling PDF models",
)
@pytest.mark.asyncio
async def test_real_docling_ocr_preserves_known_scanned_text(tmp_path):
    """Recover known setback values from a generated image-only PDF"""
    source = tmp_path / "scanned_ordinance.pdf"
    image = Image.new("RGB", (1600, 900), "white")
    drawing = ImageDraw.Draw(image)
    drawing.multiline_text(
        (80, 80),
        "EXAMPLE COUNTY WIND ORDINANCE\n\n"
        "Road setback: 125 feet\n\n"
        "Building setback: 250 feet",
        fill="black",
        font=ImageFont.load_default(size=36),
        spacing=24,
    )
    image.save(source, format="PDF", resolution=150)
    image.close()
    async with RunningAsyncServices([FileLoader()]):
        doc, raw = await read_docling_local_file(
            source,
            pytesseract_exe_fp=PYT_CMD,
            pdf_pipeline_options={
                "do_table_structure": False,
                "document_timeout": 120,
            },
        )

    assert not doc.empty
    assert doc.attrs["from_ocr"]
    assert doc.attrs["num_pages"] == 1
    assert doc.attrs["conversion_status"] == "success"
    assert "EXAMPLE COUNTY WIND ORDINANCE" in doc.text.upper()
    assert "125" in doc.text
    assert "250" in doc.text
    assert raw == source.read_bytes()


@pytest.mark.asyncio
@pytest.mark.parametrize("extension", ["md", "txt", "html"])
async def test_real_docling_local_conversion_and_cache(
    small_ordinance, extension
):
    """Preserve real headings, table values, metadata, and cached bytes"""
    source = small_ordinance[extension]
    loader = AsyncLocalDoclingFileLoader(
        doc_attrs={"source": "https://example.com/ordinance"},
        file_cache_coroutine=TempFileCache.call,
        pdf_pipeline_options={"document_timeout": 30},
    )
    async with RunningAsyncServices([FileLoader(), TempFileCache()]):
        doc = await loader.fetch(source)

        assert not doc.empty
        for expected in (
            "Example County Wind Ordinance",
            "minimum setback",
            "Roads",
            "Buildings",
            "125 feet",
            "250 feet",
        ):
            assert expected in doc.text
        assert "|" in doc.text
        assert doc.attrs["doc_type"] == (
            "html" if extension == "html" else "md"
        )
        assert doc.attrs["conversion_status"] == "success"
        assert doc.attrs["source_fp"] == source
        assert doc.attrs["source"] == "https://example.com/ordinance"
        assert not doc.attrs["from_ocr"]
        cache_fp = Path(doc.attrs["cache_fn"])
        if extension == "html":
            assert cache_fp.read_text(encoding="utf-8") == doc.text
        else:
            assert cache_fp.read_bytes() == source.read_bytes()
        assert doc.attrs["checksum"]


@pytest.mark.asyncio
async def test_real_docling_web_html_conversion(small_ordinance, monkeypatch):
    """Parse downloaded HTML with its original URI and real converter"""
    raw = small_ordinance["html"].read_bytes()
    url = "https://example.com/code/ordinance"
    loader = AsyncDoclingWebFileLoader(
        file_cache_coroutine=TempFileCache.call,
        pdf_pipeline_options={"document_timeout": 30},
    )
    monkeypatch.setattr(
        loader.content_fetcher,
        "fetch",
        AsyncMock(
            return_value=(
                raw,
                "text/html",
                "utf-8",
                {
                    "Content-Type": "text/html",
                    "Content-Disposition": 'inline; filename="ordinance.html"',
                },
            )
        ),
    )
    async with RunningAsyncServices([FileLoader(), TempFileCache()]):
        doc = await loader.fetch(url)

        assert not doc.empty
        assert "Example County Wind Ordinance" in doc.text
        assert "250 feet" in doc.text
        assert "https://example.com/code/appendix.html" in doc.text
        assert doc.attrs["source"] == url
        assert doc.attrs["doc_filename"] == "ordinance"
        assert doc.attrs["doc_type"] == "html"
        assert (
            Path(doc.attrs["cache_fn"]).read_text(encoding="utf-8") == doc.text
        )


@pytest.mark.asyncio
async def test_real_docling_invalid_pdf_returns_empty_document():
    """Handle corrupt downloaded PDFs without downloading model assets"""
    async with RunningAsyncServices([FileLoader()]):
        doc = await read_docling_web_file(
            b"not a PDF",
            "corrupt.pdf",
            pdf_pipeline_options={"document_timeout": 30},
        )

    assert doc.empty
    assert doc.attrs["doc_type"] == "unknown"


@pytest.mark.asyncio
async def test_real_docling_collection_reload_preserves_artifacts(
    small_ordinance,
):
    """Reload persisted parsed text and cache its unchanged source file"""
    doc_info = {
        "parsed_fp": small_ordinance["txt"],
        "source_fp": small_ordinance["html"],
        "source": "https://example.com/ordinance",
        "jurisdiction": "Example County",
        "from_steps": ["search"],
    }
    async with RunningAsyncServices([FileLoader(), TempFileCacheCopier()]):
        docs = await load_collected_docs(
            {"results": [doc_info]}, task_name="Example County"
        )

        assert len(docs) == 1
        doc = docs[0]
        assert not doc.empty
        assert "Example County Wind Ordinance" in doc.text
        assert "250 feet" in doc.text
        assert not doc.remove_comments
        for key, value in doc_info.items():
            assert doc.attrs[key] == value
        cache_fp = Path(doc.attrs["cache_fn"])
        assert cache_fp != small_ordinance["html"]
        assert cache_fp.read_bytes() == small_ordinance["html"].read_bytes()


@pytest.mark.asyncio
async def test_docling_collection_reload_without_parsed_path_warns():
    """Warn and skip records that cannot be reparsed from persisted text"""
    assert await load_collected_docs({}, task_name="Example County") == []
    with pytest.warns(COMPASSWarning, match="Parsed file path"):
        docs = await load_collected_docs(
            {"results": [{"source": "https://example.com/ordinance"}]},
            task_name="Example County",
        )
    assert docs == [None]


@pytest.fixture(autouse=True)
async def subprocess_log_listener():
    """Drain worker logs until process-pool shutdown completes"""
    async with LogListener([]):
        yield


@pytest.mark.skipif(
    os.getenv("GITHUB_ACTIONS") == "true", reason="Docling too heavy for GHA"
)
@pytest.mark.asyncio
async def test_basic_load_pdf(test_data_files_dir):
    """Test basic loading of local PDF document"""
    test_fp = test_data_files_dir / "Caneadea New York.pdf"

    fl = AsyncLocalDoclingFileLoader()
    async with RunningAsyncServices([FileLoader()]):
        docs = await load_docs([test_fp], fl)

    assert len(docs) == 1
    doc = docs[0]
    assert not doc.empty
    assert Path(doc.attrs.get("source_fp")) == test_fp
    assert len(doc.pages) == 1
    assert doc.attrs["num_pages"] == 3
    assert doc.attrs["doc_type"] == "pdf"
    assert "Town of Caneadea" in doc.text
    assert "Section 15.4" in doc.text


@pytest.mark.skipif(
    os.getenv("GITHUB_ACTIONS") == "true", reason="Docling too heavy for GHA"
)
@pytest.mark.asyncio
async def test_basic_load_html(test_data_files_dir):
    """Test basic loading of local HTML document"""
    test_fp = test_data_files_dir / "Whatcom.txt"

    fl = AsyncLocalDoclingFileLoader()
    async with RunningAsyncServices([FileLoader()]):
        docs = await load_docs([test_fp], fl)

    assert len(docs) == 1
    doc = docs[0]
    assert not doc.empty
    assert Path(doc.attrs.get("source_fp")) == test_fp
    assert len(doc.pages) == 1


@pytest.mark.skipif(
    os.getenv("GITHUB_ACTIONS") == "true", reason="Docling too heavy for GHA"
)
@pytest.mark.asyncio
async def test_basic_load_pdf_with_service(test_data_files_dir):
    """Test basic loading of local PDF document with service"""
    test_fp = test_data_files_dir / "Caneadea New York.pdf"

    with pytest.raises(
        COMPASSNotInitializedError,
        match=r"Must initialize the queue for 'FileLoader'.",
    ):
        await read_docling_local_file(test_fp)

    fl = AsyncLocalDoclingFileLoader()
    async with RunningAsyncServices([FileLoader()]):
        doc, __ = await read_docling_local_file(test_fp)
        doc_2 = await load_docs([test_fp], fl)

    assert not doc.empty
    assert not doc_2[0].empty
    assert doc.text == doc_2[0].text


@pytest.mark.skipif(
    os.getenv("GITHUB_ACTIONS") == "true" or not PYT_CMD,
    reason="requires PyTesseract command to be set; Docling too heavy for GHA",
)
@pytest.mark.asyncio
async def test_basic_load_ocr_pdf_with_service(test_data_files_dir):
    """Test basic loading of local PDF document with service"""
    test_fp = test_data_files_dir / "Sedgwick Kansas.pdf"

    with pytest.raises(
        COMPASSNotInitializedError,
        match=r"Must initialize the queue for 'FileLoader'.",
    ):
        await read_docling_local_file(test_fp)

    fl = AsyncLocalDoclingFileLoader(pytesseract_exe_fp=PYT_CMD)
    async with RunningAsyncServices([FileLoader()]):
        doc, __ = await read_docling_local_file(
            test_fp, pytesseract_exe_fp=PYT_CMD
        )
        doc_2 = await load_docs([test_fp], fl)

    assert not doc.empty
    assert not doc_2[0].empty
    assert doc.text == doc_2[0].text


@pytest.mark.skipif(
    os.getenv("GITHUB_ACTIONS") == "true", reason="Docling too heavy for GHA"
)
@pytest.mark.asyncio
async def test_basic_load_html_with_service(test_data_files_dir):
    """Test basic loading of local HTML document with service"""
    test_fp = test_data_files_dir / "Whatcom.txt"

    with pytest.raises(
        COMPASSNotInitializedError,
        match=r"Must initialize the queue for 'FileLoader'.",
    ):
        await read_docling_local_file(test_fp)

    fl = AsyncLocalDoclingFileLoader()
    async with RunningAsyncServices([FileLoader()]):
        doc, __ = await read_docling_local_file(test_fp)
        doc_2 = await load_docs([test_fp], fl)

    assert not doc.empty
    assert not doc_2[0].empty
    assert doc.text == doc_2[0].text


def test_docling_subprocess_supports_spawn(monkeypatch):
    """Use spawn when fork is unavailable on the current platform"""
    monkeypatch.setattr(
        multiprocessing, "get_all_start_methods", lambda: ["spawn"]
    )
    assert (
        _run_docling_in_subprocess(
            _return_from_subprocess, args=("spawned",), kwargs={}, timeout=60
        )
        == "spawned"
    )


if __name__ == "__main__":
    pytest.main(["-q", "--show-capture=all", Path(__file__), "-rapP"])
